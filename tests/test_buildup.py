from openpyxl import load_workbook

from feeprop import pricing
from feeprop.buildup import write_buildup


def test_buildup_workbook_structure(tmp_path, rate_card, analysis, estimate):
    model = pricing.price(analysis, estimate, rate_card)
    out = tmp_path / "fee_buildup.xlsx"
    write_buildup(model, rate_card, out)

    wb = load_workbook(out)
    assert wb.sheetnames == ["BUILDUP", "CLIENT VIEW"]

    ws = wb["BUILDUP"]
    # rates row mirrors the rate card (3 roles in the fixture: D, E, F)
    assert ws["C1"].value == "RATE"
    assert [ws.cell(row=1, column=4 + i).value for i in range(3)] == [300, 200, 100]
    assert [ws.cell(row=3, column=4 + i).value for i in range(3)] == ["PIC", "PM", "PA"]

    # every task row's fee is a live SUMPRODUCT against the rates row
    sumproducts = [c.value for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str) and c.value.startswith("=SUMPRODUCT")]
    assert sumproducts and all("D$1:F$1" in f for f in sumproducts)

    cv = wb["CLIENT VIEW"]
    sumifs = [c.value for row in cv.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("=SUMIF")]
    # one SUMIF per deliverable, pulling from the BUILDUP code/fee columns
    assert len(sumifs) == len(model.fee_by_deliverable)
    assert all("BUILDUP!" in f for f in sumifs)

    # the lock is stamped so an edited buildup can be traced back
    notes = [c.value for row in cv.iter_rows() for c in row
             if isinstance(c.value, str) and model.lock.lock_id in c.value]
    assert notes


def test_buildup_hours_placed_per_role(tmp_path, rate_card, analysis, estimate):
    model = pricing.price(analysis, estimate, rate_card)
    out = tmp_path / "b.xlsx"
    write_buildup(model, rate_card, out)
    ws = load_workbook(out)["BUILDUP"]
    # total hours entered across role columns == model total hours
    total = sum(
        c.value for row in ws.iter_rows(min_col=4, max_col=6) for c in row
        if isinstance(c.value, (int, float)) and c.row > 3)
    assert total == sum(model.hours_by_role.values())
