"""Render the drafted proposal (markdown) to a .docx for submission.

A deliberately small converter for the markdown feeprop itself produces:
headings, paragraphs, bold/italic runs, bullet and numbered lists, tables,
and horizontal rules. Not a general markdown engine.
"""

from __future__ import annotations

import re
from pathlib import Path

INLINE_RE = re.compile(r"(\*\*.+?\*\*|_[^_]+_)")


def _add_runs(paragraph, text: str) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _is_table_row(line: str) -> bool:
    return line.startswith("|") and line.rstrip().endswith("|")


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx(markdown: str, path: str | Path) -> None:
    try:
        import docx
        from docx.shared import Pt
    except ImportError as e:
        raise RuntimeError(
            "DOCX output requires the 'docs' extra: pip install 'feeprop[docs]'") from e

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            p = doc.add_heading("", level=min(level, 4))
            _add_runs(p, text)
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            doc.add_paragraph().add_run("").add_break()
            i += 1
            continue

        if _is_table_row(stripped):
            rows = []
            while i < len(lines) and _is_table_row(lines[i].strip()):
                cells = _split_row(lines[i].strip())
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                for r, cells in enumerate(rows):
                    for c, cell in enumerate(cells):
                        para = table.cell(r, c).paragraphs[0]
                        if r == 0:
                            para.add_run(cell.strip("*")).bold = True
                        else:
                            _add_runs(para, cell)
            continue

        if stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, stripped[2:])
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, m.group(2))
            i += 1
            continue

        # plain paragraph: join consecutive non-blank, non-structural lines
        block = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith("#") or _is_table_row(nxt)
                    or nxt.startswith(("- ", "* ")) or re.match(r"^\d+\.\s", nxt)
                    or nxt in ("---", "***", "___")):
                break
            block.append(nxt)
            i += 1
        p = doc.add_paragraph()
        _add_runs(p, " ".join(block))

    doc.save(str(path))
