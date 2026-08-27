"""Fee workbook: the xlsx the principal actually reviews before signing.

Sheets: Summary (the number, by-year, escalation capture, margins, warnings),
Detail (every hours line), Assumptions.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .models import FeeModel, ScopeAnalysis

BOLD = Font(bold=True)
H1 = Font(bold=True, size=14)
MONEY = "#,##0"
MONEY_C = "#,##0.00"


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def write_workbook(model: FeeModel, analysis: ScopeAnalysis,
                   path: str | Path) -> None:
    wb = Workbook()

    # ---- Summary -----------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{model.project_name} - Fee Summary"
    ws["A1"].font = H1
    ws["A2"] = f"Client: {model.client_name}"
    ws["A3"] = (f"Term: {model.term_months} months | Escalation: "
                f"{model.escalation_pct:g}%/yr | Lock: {model.lock.lock_id}")

    row = 5
    ws.cell(row=row, column=1, value="Fee by Contract Year").font = BOLD
    row += 1
    for y, v in sorted(model.fee_by_year.items()):
        ws.cell(row=row, column=1, value=f"Year {y}")
        ws.cell(row=row, column=2, value=v).number_format = MONEY
        row += 1
    ws.cell(row=row, column=1, value="Subtotal").font = BOLD
    ws.cell(row=row, column=2, value=model.subtotal).number_format = MONEY
    row += 1
    if model.contingency_amount:
        ws.cell(row=row, column=1, value=f"Contingency ({model.contingency_pct:g}%)")
        ws.cell(row=row, column=2, value=model.contingency_amount).number_format = MONEY
        row += 1
    if model.rounding_adjustment:
        ws.cell(row=row, column=1, value="Rounding")
        ws.cell(row=row, column=2, value=model.rounding_adjustment).number_format = MONEY
        row += 1
    ws.cell(row=row, column=1, value="TOTAL FEE").font = BOLD
    c = ws.cell(row=row, column=2, value=model.total_fee)
    c.number_format = MONEY
    c.font = BOLD
    row += 2

    ws.cell(row=row, column=1, value="Pricing Discipline").font = BOLD
    row += 1
    ws.cell(row=row, column=1, value="Fee at year-1 rates flat (the trap)")
    ws.cell(row=row, column=2, value=model.flat_rate_subtotal).number_format = MONEY
    row += 1
    ws.cell(row=row, column=1, value="Escalation capture (margin protected)")
    ws.cell(row=row, column=2, value=model.escalation_capture).number_format = MONEY
    row += 1
    if model.fee_pct_of_construction is not None:
        ws.cell(row=row, column=1, value="Fee as % of construction value")
        ws.cell(row=row, column=2, value=f"{model.fee_pct_of_construction:.2f}%")
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Margin by Role").font = BOLD
    row += 1
    for h, col in (("Role", 1), ("Revenue", 2), ("Cost", 3), ("Margin %", 4)):
        ws.cell(row=row, column=col, value=h).font = BOLD
    row += 1
    for mg in model.margins:
        ws.cell(row=row, column=1, value=mg.role_title)
        ws.cell(row=row, column=2, value=mg.revenue).number_format = MONEY
        if mg.cost is not None:
            ws.cell(row=row, column=3, value=mg.cost).number_format = MONEY
        if mg.margin_pct is not None:
            label = f"{mg.margin_pct:.1f}%" + ("  << BELOW FLOOR" if mg.below_floor else "")
            ws.cell(row=row, column=4, value=label)
        row += 1
    row += 1

    if model.warnings:
        ws.cell(row=row, column=1, value="Warnings").font = BOLD
        row += 1
        for w in model.warnings:
            cell = ws.cell(row=row, column=1, value=w)
            cell.alignment = Alignment(wrap_text=True)
            row += 1
    _autosize(ws, {1: 44, 2: 16, 3: 16, 4: 22})

    # ---- Detail ------------------------------------------------------------
    ws = wb.create_sheet("Detail")
    headers = ["Deliverable", "Role", "Phase", "Cadence", "Qty", "Hrs each",
               "Total hrs", "Base rate", "Fee (escalated)", "Fee (flat)", "Rationale"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h).font = BOLD
    for r, pl in enumerate(model.lines, start=2):
        ws.cell(row=r, column=1, value=pl.deliverable)
        ws.cell(row=r, column=2, value=pl.role_title)
        ws.cell(row=r, column=3, value=pl.phase)
        ws.cell(row=r, column=4, value=pl.cadence.value)
        ws.cell(row=r, column=5, value=pl.quantity)
        ws.cell(row=r, column=6, value=pl.hours_each)
        ws.cell(row=r, column=7, value=pl.total_hours)
        ws.cell(row=r, column=8, value=pl.base_rate).number_format = MONEY_C
        ws.cell(row=r, column=9, value=pl.amount).number_format = MONEY
        ws.cell(row=r, column=10, value=pl.flat_amount).number_format = MONEY
        ws.cell(row=r, column=11, value=pl.rationale)
    _autosize(ws, {1: 34, 2: 24, 3: 16, 4: 10, 5: 8, 6: 9, 7: 10, 8: 11, 9: 14, 10: 12, 11: 60})

    # ---- Assumptions -------------------------------------------------------
    ws = wb.create_sheet("Assumptions")
    for col, h in enumerate(["Topic", "Assumption", "Basis of fee", "Why it matters"], 1):
        ws.cell(row=1, column=col, value=h).font = BOLD
    for r, gap in enumerate(analysis.gaps, start=2):
        ws.cell(row=r, column=1, value=gap.topic)
        ws.cell(row=r, column=2, value=gap.proposed_assumption).alignment = Alignment(wrap_text=True)
        ws.cell(row=r, column=3, value=gap.pricing_basis).alignment = Alignment(wrap_text=True)
        ws.cell(row=r, column=4, value=gap.why_it_matters).alignment = Alignment(wrap_text=True)
    _autosize(ws, {1: 26, 2: 55, 3: 40, 4: 45})

    wb.save(str(path))
