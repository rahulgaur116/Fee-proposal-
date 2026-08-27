"""Read scope documents: .txt/.md natively, .pdf/.docx via optional extras."""

from __future__ import annotations

from pathlib import Path


def read_document(path: str | Path) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in (".txt", ".md", ".markdown", ""):
        return p.read_text(encoding="utf-8")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError(
                "PDF support requires the 'docs' extra: pip install 'feeprop[docs]'") from e
        reader = PdfReader(str(p))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        try:
            import docx
        except ImportError as e:
            raise RuntimeError(
                "DOCX support requires the 'docs' extra: pip install 'feeprop[docs]'") from e
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported document type: {suffix} ({p})")
