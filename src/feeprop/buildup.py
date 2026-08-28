"""Client-format fee buildup workbook.

Mirrors the firm's standard two-sheet fee builder:

  BUILDUP     - rates row across the top, one row per task with hours by role,
                fee per row as =SUMPRODUCT(rates, hours), a summary code per
                deliverable, block subtotals with a blended-rate check.
  CLIENT VIEW - one row per client-facing deliverable, pulling its lump sum
                from the buildup via =SUMIF on the summary code.

All money cells are live formulas so the workbook recalculates when hours are
edited - but the shipped numbers must match the locked fee model, so callers
should verify totals after recalculation.
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .models import FeeModel, RateCard

FONT = "Arial"
BOLD = Font(name=FONT, bold=True)
PLAIN = Font(name=FONT)
MONEY = "$#,##0"
RATE_FMT = "#,##0"


def _codes(model: FeeModel) -> dict[str, str]:
    """Summary code letter per deliverable, in order of first appearance."""
    out: dict[str, str] = {}
    for pl in model.lines:
        if pl.deliverable not in out:
            out[pl.deliverable] = chr(ord("A") + len(out))
    return out


def write_buildup(model: FeeModel, rate_card: RateCard, path: str | Path) -> None:
    roles = rate_card.roles
    n = len(roles)
    ROLE0 = 4                    # D
    REIMB = ROLE0 + n + 1        # skip one column after the roles
    FEE = REIMB + 2              # skip one more
    CODE = FEE + 1
    fee_l, code_l = get_column_letter(FEE), get_column_letter(CODE)
    role_ls = [get_column_letter(ROLE0 + i) for i in range(n)]
    rate_range = f"{role_ls[0]}$1:{role_ls[-1]}$1"

    codes = _codes(model)

    wb = Workbook()

    # ---- BUILDUP ----------------------------------------------------------
    ws = wb.active
    ws.title = "BUILDUP"
    ws.cell(row=1, column=3, value="RATE").font = BOLD
    for i, r in enumerate(roles):
        c = ws.cell(row=1, column=ROLE0 + i, value=r.bill_rate)
        c.font = BOLD
        c.number_format = RATE_FMT
    ws.cell(row=2, column=ROLE0, value="HOURS").font = BOLD
    ws.cell(row=2, column=REIMB, value="REIMBURSABLES").font = BOLD
    ws.cell(row=2, column=FEE, value="FEE").font = BOLD
    ws.cell(row=2, column=CODE, value="SUMMARY CODE").font = BOLD
    ws.cell(row=3, column=2, value="Comments & Notes").font = BOLD
    ws.cell(row=3, column=3, value="Task").font = BOLD
    for i, r in enumerate(roles):
        ws.cell(row=3, column=ROLE0 + i, value=r.code).font = BOLD
    ws.cell(row=3, column=REIMB, value="LUMP SUM").font = BOLD
    ws.cell(row=3, column=FEE, value="TOTAL").font = BOLD

    # group lines into (deliverable -> task -> {role: hours})
    grouped: dict[str, dict[str, dict[str, float]]] = defaultdict(dict)
    for pl in model.lines:
        task = grouped[pl.deliverable].setdefault(pl.rationale, {})
        task[pl.role_code] = task.get(pl.role_code, 0.0) + pl.total_hours

    row = 4
    block_subtotal_rows: list[int] = []
    first_task_row = row
    for deliverable, tasks in grouped.items():
        code = codes[deliverable]
        head = ws.cell(row=row, column=3, value=f"{deliverable}  [{code}]")
        head.font = BOLD
        row += 1
        block_start = row
        for task, hours in tasks.items():
            cell = ws.cell(row=row, column=3, value=task)
            cell.font = PLAIN
            cell.alignment = Alignment(wrap_text=True)
            for i, r in enumerate(roles):
                if r.code in hours:
                    ws.cell(row=row, column=ROLE0 + i, value=hours[r.code])
            f = ws.cell(row=row, column=FEE,
                        value=f"=SUMPRODUCT({rate_range},"
                              f"{role_ls[0]}{row}:{role_ls[-1]}{row})")
            f.number_format = MONEY
            ws.cell(row=row, column=CODE, value=code)
            row += 1
        # block subtotal
        for i in range(n):
            col_l = role_ls[i]
            ws.cell(row=row, column=ROLE0 + i,
                    value=f"=SUBTOTAL(9,{col_l}{block_start}:{col_l}{row - 1})").font = BOLD
        sf = ws.cell(row=row, column=FEE,
                     value=f"=SUBTOTAL(9,{fee_l}{block_start}:{fee_l}{row - 1})")
        sf.font = BOLD
        sf.number_format = MONEY
        hrs_col = get_column_letter(CODE + 1)
        ws.cell(row=row, column=CODE + 1,
                value=f"=SUM({role_ls[0]}{row}:{role_ls[-1]}{row})").font = BOLD
        blended = ws.cell(
            row=row, column=CODE + 2,
            value=f"=IF({hrs_col}{row}=0,\"\",{fee_l}{row}/{hrs_col}{row})")
        blended.number_format = MONEY
        ws.cell(row=row, column=3, value="Subtotal").font = BOLD
        block_subtotal_rows.append(row)
        row += 2

    # grand total
    ws.cell(row=row, column=3, value="GRAND TOTAL").font = BOLD
    for i in range(n):
        col_l = role_ls[i]
        ws.cell(row=row, column=ROLE0 + i,
                value=f"=SUBTOTAL(9,{col_l}{first_task_row}:{col_l}{row - 1})").font = BOLD
    gt = ws.cell(row=row, column=FEE,
                 value=f"=SUBTOTAL(9,{fee_l}{first_task_row}:{fee_l}{row - 1})")
    gt.font = BOLD
    gt.number_format = MONEY

    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 70
    for i in range(n):
        ws.column_dimensions[role_ls[i]].width = 7
    ws.column_dimensions[get_column_letter(REIMB)].width = 15
    ws.column_dimensions[fee_l].width = 12
    ws.column_dimensions[code_l].width = 15

    # ---- CLIENT VIEW ------------------------------------------------------
    cv = wb.create_sheet("CLIENT VIEW")
    cv.cell(row=2, column=2, value="ESTIMATING SCOPE OF SERVICES").font = BOLD
    cv.cell(row=2, column=7, value="SUMMARY CODE").font = BOLD
    cv.cell(row=3, column=2, value="Project").font = BOLD
    cv.cell(row=3, column=3, value="Task").font = BOLD
    cv.cell(row=3, column=4, value="Lump Sum Fee").font = BOLD
    cv.cell(row=3, column=5, value="Deliverable").font = BOLD
    r = 5
    cv.cell(row=r, column=2, value=model.project_name).font = PLAIN
    first_cv_row = r
    for deliverable, code in codes.items():
        cv.cell(row=r, column=3, value=deliverable).font = PLAIN
        f = cv.cell(row=r, column=4,
                    value=f"=SUMIF(BUILDUP!{code_l}:{code_l},G{r},"
                          f"BUILDUP!{fee_l}:{fee_l})")
        f.number_format = MONEY
        cv.cell(row=r, column=7, value=code)
        r += 1
    cv.cell(row=r, column=3, value="TOTAL LUMP SUM FEE").font = BOLD
    tf = cv.cell(row=r, column=4, value=f"=SUBTOTAL(9,D{first_cv_row}:D{r - 1})")
    tf.font = BOLD
    tf.number_format = MONEY
    r += 2
    cv.cell(row=r, column=2,
            value=f"Locked fee model {model.lock.lock_id}: total "
                  f"${model.total_fee:,.0f}. If the buildup is edited, "
                  f"re-run `feeprop price` so the proposal and this workbook "
                  f"stay on one number.").font = PLAIN

    cv.column_dimensions["B"].width = 40
    cv.column_dimensions["C"].width = 55
    cv.column_dimensions["D"].width = 15
    cv.column_dimensions["E"].width = 50
    cv.column_dimensions["G"].width = 15

    wb.save(str(path))
